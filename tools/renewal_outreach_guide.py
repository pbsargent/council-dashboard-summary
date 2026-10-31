"""Reader-guide supplement shared by the maintained guide generator and updates."""
from docx.shared import Pt, RGBColor

PARAGRAPHS = (
    ('Daily Movo renewal cohorts', 'Open District Performance > Renewal Outreach to review members due, renewed, renewal rate and outstanding membership renewals. The Overview also links to the page. These are source-reported member aggregates; person-level deduplication and concurrent registrations cannot be independently verified from this summary.'),
    ('Filters and comparisons', 'Cohort and District apply to renewal cards, district comparisons, adult/youth progress, reported decline reasons, daily history and the filtered CSV. This source has no program-type breakdown. Movo outreach and renewal timing cover both cohorts and follow District only; that scope is stated on the page.'),
    ('Renewal and follow-up measures', 'Renewal rate = renewed divided by members due. Outstanding = members due minus renewed, including said no. A zero denominator is unavailable. Adult/youth renewals and reason breakdowns remain unavailable where not supplied. Not reached is a source renewal state and differs from not yet messaged. Reply rate = replied divided by members messaged. Renewal timing before or after the first message does not establish causal impact. These measures are not annual retention.'),
    ('Freshness and daily history', 'Source updated is the daily Google Sheet modification time in America/Chicago. The source has no separate Movo observation or report-generation timestamp. A stale notice appears when the source update date is earlier than today. Daily refresh requires a complete same-day source capture, reconciled district/council counts and valid cohort dates. Failures retain the last validated publication. History begins with the first validated capture, keeps the latest update per source day and shows up to 90 days; missing days are not filled in.'),
    ('Public data', 'Only council and district aggregates are published. No member names, member IDs, contact details, private source paths or credentials are included. Use the live Dashboard Guide for current page controls and source behavior.'),
)


def append_renewal_outreach(doc):
    doc.add_page_break()
    heading = doc.add_paragraph()
    heading.paragraph_format.space_after = Pt(12)
    heading.paragraph_format.keep_with_next = True
    run = heading.add_run('Renewal Outreach')
    run.font.name = 'Arial'
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string('051F4C')
    for title, text in PARAGRAPHS:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(9)
        paragraph.paragraph_format.line_spacing = 1.08
        for value, bold in ((title + '\n', True), (text, False)):
            run = paragraph.add_run(value)
            run.font.name = 'Arial'
            run.font.size = Pt(10)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string('051F4C')

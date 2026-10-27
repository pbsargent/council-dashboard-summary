"""Publish only validated unit aggregates from the newest dated joins workbooks."""
import argparse
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re
import tempfile
from zoneinfo import ZoneInfo

import openpyxl

DEFAULT_SOURCE = Path('/Users/petersargent/Library/CloudStorage/GoogleDrive-petersargent@capitolareascouting.org/Shared drives/CAC-Published-Reports')
KINDS = {'cubs': ('Cub_Pack', 'Pack', 'Pack joins'), 'scouts-bsa': ('Scout_Troop', 'Troop', 'Troop joins')}


def require(ok, message):
    if not ok:
        raise ValueError(message)


def count(value):
    require(type(value) in (int, float) and value >= 0 and int(value) == value, f'Invalid count: {value!r}')
    return int(value)


def newest(folder, kind):
    pattern = re.compile(r'CAC_All_' + KINDS[kind][0] + r'_Joins_Sept1_2026_to_(\d{4}-\d{2}-\d{2})(?:_(\d{6})_(?:CDT|CST))?\.xlsx')
    candidates = [(date.fromisoformat(m[1]), datetime.strptime(m[2] or '000000', '%H%M%S').time(), p) for p in folder.iterdir() if (m := pattern.fullmatch(p.name))]
    require(candidates, f'Missing {kind} joins workbook in source folder')
    latest = max(candidates)
    return latest[0], latest[2]


def read_report(path, report_date, kind):
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    _, unit_type, main_sheet = KINDS[kind]
    require(main_sheet in workbook.sheetnames, f'Missing {main_sheet}')

    def parse(sheet):
        rows = list(sheet.values)
        require(len(rows) >= 8, 'Incomplete joins sheet')
        header = list(rows[7])
        adult_schema = 'New adults to unit' in header
        metric_headers = ['New youth joins', 'New adults to unit', 'New to Scouting (reported)'] if adult_schema else ['New joins']
        expected = ['Rank' if sheet.title == main_sheet else 'Council rank', unit_type, 'District', *metric_headers, 'Unit ID', 'Chartered organization']
        require(header[:len(expected)] == expected, 'Unexpected unit schema')
        prefix = f'September 1, 2026–{report_date.strftime("%B")} {report_date.day}, {report_date.year}. '
        wording = 'Youth: New only. Adults: distinct people, new to unit / reported New to Scouting.' if adult_schema else 'New youth only. Transfers and re-registrations excluded.'
        require(rows[2][0] == prefix + wording, 'Workbook period does not match filename')
        age_start = len(expected)
        ages = header[age_start:age_start+9]
        require(ages == [f'Age {n}' for n in (range(4, 12) if kind == 'cubs' else range(10, 18))] + ['Age unavailable'], 'Unexpected age schema')
        review_index = age_start + 9
        require(header[review_index] == 'Review status', 'Unexpected review schema')
        units = []
        for row in rows[8:]:
            if not any(v is not None for v in row[:6]):
                continue
            require(len(row) > review_index, 'Incomplete unit row')
            rank, unit, district, joins = row[:4]
            unit_id, organization = row[age_start-2:age_start]
            require(isinstance(unit, str) and unit.startswith(unit_type + ' '), 'Invalid unit label')
            require(isinstance(district, str) and district.strip(), 'Missing district')
            require(str(unit_id).isdigit() or (unit_id is None and bool(row[review_index])), 'Invalid unit identity')
            age_counts = [count(v) for v in row[age_start:review_index]]
            require(sum(age_counts) == count(joins), f'Age totals do not reconcile for {unit_id}')
            require(row[review_index] is None or isinstance(row[review_index], str), 'Invalid review status')
            units.append(dict(rank=count(rank), unit=unit, district=district, joins=count(joins), adults_new_to_unit=count(row[4]) if adult_schema else None, adults_new_to_scouting_reported=count(row[5]) if adult_schema else None, unit_id=str(unit_id) if unit_id is not None else None, organization=organization, ages=age_counts, review=row[review_index] or ''))
        ids = [u['unit_id'] for u in units if u['unit_id'] is not None]
        require(len(set(ids)) == len(ids), 'Duplicate unit IDs')
        require(len({(u['district'], u['unit']) for u in units}) == len(units), 'Duplicate unit labels')
        require(len(units) == count(rows[4][1]), 'Unit count does not reconcile')
        require(sum(u['joins'] for u in units) == count(rows[5][1]), 'Join total does not reconcile')
        require([sum(u['ages'][i] for u in units) for i in range(9)] == [count(v) for v in rows[5][age_start:review_index]], 'Summary ages do not reconcile')
        if adult_schema:
            require(list(rows[4][3:6]) == metric_headers, 'Missing adult summary headers')
            for index, field in ((3,'joins'), (4,'adults_new_to_unit'), (5,'adults_new_to_scouting_reported')):
                require(sum(u[field] for u in units) == count(rows[5][index]), f'{field} summary does not reconcile')
        review = re.match(r'(\d+) unit\(s\) need review', rows[3][0] or '')
        require(review and int(review[1]) == sum(bool(u['review']) for u in units), 'Review coverage does not reconcile')
        return units, ages

    units, ages = parse(workbook[main_sheet])
    districts = {u['district'] for u in units}
    require(set(workbook.sheetnames) == districts | {main_sheet}, 'District sheet coverage does not reconcile')
    for district in districts:
        detail, detail_ages = parse(workbook[district])
        require(detail_ages == ages and detail == [u for u in units if u['district'] == district], f'District detail differs: {district}')
    workbook.close()
    require(before == hashlib.sha256(path.read_bytes()).hexdigest(), 'Workbook changed while reading')
    return dict(program=kind, report_date=report_date.isoformat(), period_start='2026-09-01', source_name=path.name, source_sha256=before, age_labels=ages, units=units)


def district_key(value):
    return re.sub(r'\s+\d{2}$', '', value.strip()).casefold()


def service_area_mapping(snapshot_path, reports, expected_date):
    snapshot = json.loads(snapshot_path.read_text())
    require(snapshot['dashboard']['report_date'] == expected_date.isoformat(), 'Service-area snapshot must match joins report date')
    mapping = {}
    for row in snapshot['dashboard']['districts']:
        key = district_key(row['district'])
        area = row['service_area']
        require(key not in mapping and isinstance(area, str) and area.strip(), 'Missing or ambiguous district service area')
        mapping[key] = area
    result = {}
    for report in reports.values():
        for district in {u['district'] for u in report['units']}:
            require(district_key(district) in mapping, f'No authoritative service area for {district}')
            result[district] = mapping[district_key(district)]
    return result


def build(folder, output, expected_date, council_snapshot=None):
    selected = {kind: newest(folder, kind) for kind in KINDS}
    require(all(day == expected_date for day, _ in selected.values()), 'Both newest joins workbooks must match the run date')
    reports = {kind: read_report(path, day, kind) for kind, (day, path) in selected.items()}
    require(all(u[field] is not None for r in reports.values() for u in r['units'] for field in ('adults_new_to_unit','adults_new_to_scouting_reported')), 'Current joins publication requires both adult measures; legacy youth-only workbooks cannot replace it')
    require(council_snapshot is not None, 'A same-day council snapshot is required for service areas')
    snapshot_before = hashlib.sha256(council_snapshot.read_bytes()).hexdigest()
    areas = service_area_mapping(council_snapshot, reports, expected_date)
    for report in reports.values():
        report['district_service_areas'] = {d: areas[d] for d in sorted({u['district'] for u in report['units']})}
    require(snapshot_before == hashlib.sha256(council_snapshot.read_bytes()).hexdigest(), 'Service-area source changed while reading')
    # Validate the complete pair before replacing the single public bundle.
    for kind, (_, path) in selected.items():
        require(newest(folder, kind) == selected[kind], 'A newer joins generation appeared during the build; retry the complete pair')
        require(hashlib.sha256(path.read_bytes()).hexdigest() == reports[kind]['source_sha256'], 'Source changed across pair read')
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(mode='w', dir=output.parent, prefix='.joins-', suffix='.js', delete=False) as handle:
        temporary = Path(handle.name)
        handle.write('window.CAC_JOINS = ' + json.dumps(reports, ensure_ascii=False, allow_nan=False) + ';\n')
    try:
        validate_public_bundle(temporary)
        for kind, (_, path) in selected.items():
            require(newest(folder, kind) == selected[kind] and hashlib.sha256(path.read_bytes()).hexdigest() == reports[kind]['source_sha256'], 'Joins sources changed before installation; retry the complete pair')
        require(snapshot_before == hashlib.sha256(council_snapshot.read_bytes()).hexdigest(), 'Service-area source changed before installation')
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
    print(json.dumps({k: {'units': len(v['units']), 'joins': sum(u['joins'] for u in v['units']), 'source': v['source_name']} for k, v in reports.items()}))


def validate_public_bundle(path):
    text = Path(path).read_text()
    require(text.startswith('window.CAC_JOINS = ') and text.endswith(';\n'), 'Invalid joins bundle wrapper')
    reports = json.loads(text[len('window.CAC_JOINS = '):-2])
    require(isinstance(reports, dict), 'Joins bundle must be an object')
    require(set(reports) == set(KINDS), 'Incomplete joins program pair')
    require(all(isinstance(r,dict) and isinstance(r.get('report_date'),str) for r in reports.values()), 'Invalid report metadata')
    require(len({r['report_date'] for r in reports.values()}) == 1, 'Mismatched joins dates')
    fields = {'rank','unit','district','joins','unit_id','organization','ages','review','adults_new_to_unit','adults_new_to_scouting_reported'}
    for kind, report in reports.items():
        require(set(report) == {'program','report_date','period_start','source_name','source_sha256','age_labels','units','district_service_areas'}, 'Unexpected public report field')
        require(report['program'] == kind and isinstance(report['units'],list) and report['units'], 'Missing joins population')
        require(isinstance(report['source_name'],str) and isinstance(report['source_sha256'],str), 'Invalid source metadata')
        require(report['period_start'] == '2026-09-01' and date.fromisoformat(report['report_date']) >= date(2026,9,1), 'Invalid joins period')
        pattern = r'CAC_All_' + KINDS[kind][0] + r'_Joins_Sept1_2026_to_' + report['report_date'] + r'(?:_\d{6}_(?:CDT|CST))?\.xlsx'
        require(re.fullmatch(pattern, report['source_name']) is not None, 'Invalid joins source identity')
        require(re.fullmatch(r'[a-f0-9]{64}', report['source_sha256']) is not None, 'Invalid joins source digest')
        require(report['age_labels'] == [f'Age {n}' for n in (range(4,12) if kind == 'cubs' else range(10,18))] + ['Age unavailable'], 'Invalid age buckets')
        require(set(report['district_service_areas']) == {u['district'] for u in report['units']}, 'Incomplete joins service-area coverage')
        require(all(isinstance(a, str) and a.strip() for a in report['district_service_areas'].values()), 'Invalid joins service area')
        for unit in report['units']:
            require(isinstance(unit,dict), 'Invalid public unit record')
            require(set(unit) == fields, 'Unexpected public unit field')
            require(isinstance(unit['unit'], str) and unit['unit'].startswith(KINDS[kind][1] + ' '), 'Invalid public unit label')
            require(isinstance(unit['district'], str) and bool(unit['district'].strip()), 'Invalid public district')
            require(unit['organization'] is None or isinstance(unit['organization'], str), 'Invalid public organization')
            require(unit['review'] in ('', 'Review — missing from current unit report', 'Review — missing from current unit report; district from join records'), 'Unexpected public review status')
            require((isinstance(unit['unit_id'],str) and unit['unit_id'].isdigit()) or (unit['unit_id'] is None and bool(unit['review'])), 'Invalid public unit identity')
            count(unit['rank'])
            require(len(unit['ages']) == 9 and sum(count(n) for n in unit['ages']) == count(unit['joins']), 'Invalid public joins counts')
            for field in ('adults_new_to_unit','adults_new_to_scouting_reported'):
                count(unit[field])
        ids = [u['unit_id'] for u in report['units'] if u['unit_id'] is not None]
        require(len(ids) == len(set(ids)), 'Duplicate public unit IDs')
        require(len({(u['district'],u['unit']) for u in report['units']}) == len(report['units']), 'Duplicate public unit labels')
    return reports


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir', type=Path, default=DEFAULT_SOURCE)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--council-snapshot', type=Path, required=True)
    parser.add_argument('--date', type=date.fromisoformat, default=datetime.now(ZoneInfo('America/Chicago')).date())
    args = parser.parse_args()
    build(args.source_dir, args.output, args.date, args.council_snapshot)
